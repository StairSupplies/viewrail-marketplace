"""Build the A3 Word template (templates/a3_template.docx) programmatically via python-docx.

Run this script once whenever the visual design changes; the resulting .docx is the master
template that docxtpl renders against. All Jinja tags ({{ var }} and {%tr ... %} loops) are
embedded as text content so docxtpl can substitute on render.

Why programmatic rather than hand-authored Word: the engineer agent has no Word UI. The
output is a normal docxtpl-compatible .docx; Garrett can open it in Word and tune visuals
afterward without breaking the Jinja tags.

Usage:
    python3 templates/build_template.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


# Brand colors lifted from the source HTML reference.
NAVY = RGBColor(0x0D, 0x21, 0x37)         # dark navy header / footer band
GOLD = RGBColor(0xE8, 0xA8, 0x38)         # accent
BLUE_LEFT = RGBColor(0x2E, 0x6D, 0xA4)    # left-column section header bar
BLUE_RIGHT = RGBColor(0x1F, 0x50, 0x80)   # right-column section header bar
TINT = RGBColor(0xEB, 0xF4, 0xFB)         # alternating row tint
TBL_HEAD = RGBColor(0x1A, 0x3F, 0x5C)     # table header rows
TBL_HEAD_TEXT = RGBColor(0xA8, 0xC8, 0xE0)
LABEL = RGBColor(0x6A, 0x8F, 0xAA)
BODY_TEXT = RGBColor(0x1A, 0x2E, 0x40)


def shade_cell(cell, hex_color: str) -> None:
    """Apply a solid background fill to a table cell (python-docx has no native shading API)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_borders(cell, color="CCE0F0", size="4") -> None:
    """Set thin light-blue borders on all 4 sides of a cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), size)
        b.set(qn("w:color"), color)
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def add_run(paragraph, text: str, *, bold=False, color: RGBColor | None = None,
            size: int | None = None, caps=False) -> None:
    """Add a styled run to a paragraph."""
    run = paragraph.add_run(text)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if size is not None:
        run.font.size = Pt(size)
    if caps:
        run.font.all_caps = True


def section_header(cell, number: int, title: str, desc: str, color_hex: str) -> None:
    """Render a section header bar (number circle + title + description) inside `cell`."""
    shade_cell(cell, color_hex)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, f"{number}  ", bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=11)
    add_run(p, title.upper(), bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    add_run(p2, desc, color=RGBColor(0xA8, 0xC8, 0xE0), size=7)


def labeled_paragraph(cell, label: str, jinja: str, *, label_color=LABEL,
                      body_color=BODY_TEXT) -> None:
    """Render a label + Jinja value pair as one paragraph inside `cell`."""
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    add_run(p, f"{label}: ", bold=True, color=label_color, size=8, caps=True)
    add_run(p, jinja, color=body_color, size=10)


def freeform_paragraph(cell, jinja: str) -> None:
    """A single-line Jinja substitution paragraph for narrative fields."""
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    add_run(p, jinja, color=BODY_TEXT, size=10)


def add_section_inner_table(cell, headers: list[str], row_jinja: list[str],
                            loop_open: str, loop_close: str = "{%tr endfor %}") -> None:
    """Add a docxtpl-friendly inner table with a single Jinja tr-loop inside `cell`.

    Structure: 4 rows.
      Row 0: header row (kept as-is on render).
      Row 1: control row containing only {%tr for ... %} (removed on render).
      Row 2: body row containing the data Jinja per cell (replicated for each item).
      Row 3: control row containing only {%tr endfor %} (removed on render).
    """
    tbl = cell.add_table(rows=4, cols=len(headers))
    tbl.autofit = True
    # Header row (row 0).
    for i, h in enumerate(headers):
        c = tbl.cell(0, i)
        shade_cell(c, "1A3F5C")
        set_cell_borders(c)
        c.paragraphs[0].paragraph_format.space_before = Pt(0)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        add_run(c.paragraphs[0], h, bold=True, color=TBL_HEAD_TEXT, size=7, caps=True)
    # Open control row (row 1): only the {%tr for%} tag.
    open_cell = tbl.cell(1, 0)
    add_run(open_cell.paragraphs[0], loop_open, color=BODY_TEXT, size=9)
    # Body row (row 2): one cell per data Jinja expression.
    for i, expr in enumerate(row_jinja):
        c = tbl.cell(2, i)
        set_cell_borders(c)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        add_run(p, expr, color=BODY_TEXT, size=9)
    # Close control row (row 3): only the {%tr endfor%} tag.
    close_cell = tbl.cell(3, 0)
    add_run(close_cell.paragraphs[0], loop_close, color=BODY_TEXT, size=9)


def build() -> Path:
    """Build the A3 template DOCX and return its path."""
    doc = Document()

    # Page setup: landscape A3.
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(42.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.8)
    section.bottom_margin = Cm(0.8)
    section.left_margin = Cm(0.8)
    section.right_margin = Cm(0.8)

    # ---------- Header band ----------
    header_tbl = doc.add_table(rows=1, cols=5)
    header_tbl.autofit = True
    title_cell = header_tbl.cell(0, 0)
    shade_cell(title_cell, "0D2137")
    p = title_cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "{{ header.title }}", bold=True,
            color=RGBColor(0xFF, 0xFF, 0xFF), size=14, caps=True)

    for idx, (label, jinja) in enumerate([
        ("Owner", "{{ header.owner }}"),
        ("Team / Area", "{{ header.team_or_area }}"),
        ("Date", "{{ header.date }}"),
        ("Revision", "{{ header.revision }}"),
    ], start=1):
        c = header_tbl.cell(0, idx)
        shade_cell(c, "0D2137")
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        add_run(p, label.upper(), bold=True,
                color=RGBColor(0x4A, 0x6A, 0x82), size=7, caps=True)
        p2 = c.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        add_run(p2, jinja, color=RGBColor(0xC8, 0xDF, 0xF0), size=9)

    # Gold accent line below the header.
    gold_p = doc.add_paragraph()
    gold_p.paragraph_format.space_before = Pt(0)
    gold_p.paragraph_format.space_after = Pt(2)
    p_pr = gold_p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), "E8A838")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    # ---------- Body: 2-column outer table ----------
    body = doc.add_table(rows=1, cols=2)
    body.autofit = False
    left = body.cell(0, 0)
    right = body.cell(0, 1)
    left.width = Cm(20)
    right.width = Cm(20)

    build_left_column(left)
    build_right_column(right)

    # ---------- Footer band ----------
    foot_tbl = doc.add_table(rows=1, cols=3)
    foot_tbl.autofit = True
    for idx, (label, jinja) in enumerate([
        ("Approved By", "{{ footer.approved_by }}"),
        ("Review Date", "{{ footer.review_date }}"),
        ("Status", "{{ footer.status }}"),
    ]):
        c = foot_tbl.cell(0, idx)
        shade_cell(c, "0D2137")
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        add_run(p, label.upper(), bold=True,
                color=RGBColor(0x3A, 0x5A, 0x72), size=7, caps=True)
        p2 = c.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        add_run(p2, jinja, color=RGBColor(0x6A, 0x8F, 0xAA), size=9)

    out = Path(__file__).parent / "a3_template.docx"
    doc.save(out)
    return out


def build_left_column(cell) -> None:
    """Sections 1 to 4 in the left column."""
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)

    # 1 Background
    inner = cell.add_table(rows=2, cols=1)
    section_header(inner.cell(0, 0), 1, "Background",
                   "Why this matters; what constraint it creates in the flow",
                   "2E6DA4")
    body = inner.cell(1, 0)
    shade_cell(body, "EBF4FB")
    body.paragraphs[0].paragraph_format.space_after = Pt(0)
    freeform_paragraph(body, "{{ section_1_background.narrative }}")
    labeled_paragraph(body, "Business impact", "{{ section_1_background.business_impact }}")
    labeled_paragraph(body, "First observed", "{{ section_1_background.first_observed }}")
    labeled_paragraph(body, "Frequency / scope", "{{ section_1_background.frequency_or_scope }}")
    labeled_paragraph(body, "Constraint created", "{{ section_1_background.constraint_created }}")

    # 2 Current Condition
    inner = cell.add_table(rows=2, cols=1)
    section_header(inner.cell(0, 0), 2, "Current Condition",
                   "What IS happening now, baseline metrics here",
                   "2E6DA4")
    body = inner.cell(1, 0)
    body.paragraphs[0].paragraph_format.space_after = Pt(0)
    freeform_paragraph(body, "{{ section_2_current_condition.narrative }}")
    # Baseline metrics: single Jinja substitution; render_docx.py joins the list.
    p = body.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "BASELINE METRICS", bold=True, color=LABEL, size=8, caps=True)
    freeform_paragraph(body, "{{ section_2_current_condition.baseline_metrics_display }}")
    labeled_paragraph(body, "Process observation",
                      "{{ section_2_current_condition.process_observation }}")
    labeled_paragraph(body, "Where in flow",
                      "{{ section_2_current_condition.where_in_flow }}")

    # 3 Goal / Target
    inner = cell.add_table(rows=2, cols=1)
    section_header(inner.cell(0, 0), 3, "Goal / Target Condition",
                   "Desired future state, target metrics and value expected",
                   "2E6DA4")
    body = inner.cell(1, 0)
    shade_cell(body, "EBF4FB")
    body.paragraphs[0].paragraph_format.space_after = Pt(0)
    add_section_inner_table(
        body,
        headers=["Metric", "Current", "Target"],
        row_jinja=["{{ m.metric }}", "{{ m.current }}", "{{ m.target }}"],
        loop_open="{%tr for m in section_3_goal_target.metrics %}",
    )
    labeled_paragraph(body, "Value expected", "{{ section_3_goal_target.value_expected }}")
    labeled_paragraph(body, "Target date", "{{ section_3_goal_target.target_date }}")

    # 4 Root Cause
    inner = cell.add_table(rows=2, cols=1)
    section_header(inner.cell(0, 0), 4, "Root Cause Analysis",
                   "5-Why or fishbone, trace each symptom to its true root cause",
                   "2E6DA4")
    body = inner.cell(1, 0)
    body.paragraphs[0].paragraph_format.space_after = Pt(0)
    labeled_paragraph(body, "Problem statement",
                      "{{ section_4_root_cause.problem_statement }}")
    p = body.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "5 WHYS", bold=True, color=LABEL, size=8, caps=True)
    freeform_paragraph(body, "{{ section_4_root_cause.five_whys_display }}")
    labeled_paragraph(body, "Confirmed by",
                      "{{ section_4_root_cause.root_cause_confirmed_by }}")
    p = body.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "ROOT CAUSES IDENTIFIED", bold=True, color=LABEL, size=8, caps=True)
    freeform_paragraph(body, "{{ section_4_root_cause.root_causes_identified_display }}")


def build_right_column(cell) -> None:
    """Sections 5 to 7 in the right column."""
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)

    # 5 Countermeasures
    inner = cell.add_table(rows=2, cols=1)
    section_header(inner.cell(0, 0), 5, "Countermeasures",
                   "Actions tied to specific root causes; each must have an owner and due date",
                   "1F5080")
    body = inner.cell(1, 0)
    body.paragraphs[0].paragraph_format.space_after = Pt(0)
    add_section_inner_table(
        body,
        headers=["Action", "Detail / Rationale", "Owner", "Due", "Status"],
        row_jinja=["{{ c.action }}", "{{ c.detail_rationale }}",
                   "{{ c.owner }}", "{{ c.due }}", "{{ c.status }}"],
        loop_open="{%tr for c in section_5_countermeasures.actions %}",
    )

    # 6 Effect Confirmation
    inner = cell.add_table(rows=2, cols=1)
    section_header(inner.cell(0, 0), 6, "Effect Confirmation / Cost-Benefit",
                   "Before vs after, quantify the expected impact",
                   "1F5080")
    body = inner.cell(1, 0)
    shade_cell(body, "EBF4FB")
    body.paragraphs[0].paragraph_format.space_after = Pt(0)
    add_section_inner_table(
        body,
        headers=["Category", "Before", "After (Expected)"],
        row_jinja=["{{ e.category }}", "{{ e.before }}", "{{ e.after_expected }}"],
        loop_open="{%tr for e in section_6_effect_confirmation.comparisons %}",
    )
    labeled_paragraph(body, "Key insight",
                      "{{ section_6_effect_confirmation.key_insight }}")

    # 7 Follow-Up
    inner = cell.add_table(rows=2, cols=1)
    section_header(inner.cell(0, 0), 7, "Follow-Up Actions",
                   "Progress metrics vs baseline; confirm countermeasures worked",
                   "1F5080")
    body = inner.cell(1, 0)
    body.paragraphs[0].paragraph_format.space_after = Pt(0)
    add_section_inner_table(
        body,
        headers=["Item", "Measure", "Owner", "Due", "Status"],
        row_jinja=["{{ f.item }}", "{{ f.measure }}",
                   "{{ f.owner }}", "{{ f.due }}", "{{ f.status }}"],
        loop_open="{%tr for f in section_7_follow_up.actions %}",
    )
    labeled_paragraph(body, "Standardize / Next steps",
                      "{{ section_7_follow_up.standardize_next_steps }}")


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
